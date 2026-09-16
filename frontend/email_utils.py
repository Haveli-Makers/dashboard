import datetime
import io
import re
import smtplib
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pandas as pd

from CONFIG import (
    SMTP_FROM_EMAIL,
    SMTP_HOST,
    SMTP_PASSWORD,
    SMTP_PORT,
    SMTP_TIMEOUT,
    SMTP_USE_TLS,
    SMTP_USERNAME,
)
from CONFIG import SAMPLES_EMAIL_BODY_TEMPLATE as SAMPLES_EMAIL_BODY_TEMPLATE
from CONFIG import SAMPLES_EMAIL_SUBJECT_TEMPLATE as SAMPLES_EMAIL_SUBJECT_TEMPLATE
from CONFIG import SPREAD_EMAIL_BODY_TEMPLATE as SPREAD_EMAIL_BODY_TEMPLATE
from CONFIG import SPREAD_EMAIL_SUBJECT_TEMPLATE as SPREAD_EMAIL_SUBJECT_TEMPLATE

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

class _SafeDict(dict):
    def __missing__(self, key):
        return "{" + key + "}"


def render_template(template: str, context: dict) -> str:
    """Render a {var}-style template, leaving unknown placeholders untouched."""
    return template.format_map(_SafeDict(context))


def build_spread_email_context(
    connectors: list[str],
    pairs: list[str],
    window_hours,
    row_count: int,
    failed_count: int,
) -> dict:
    """Build the variable context used to render the spread report email template."""
    now = datetime.datetime.now()
    return {
        "connectors": ", ".join(connectors) if connectors else "-",
        "pairs": ", ".join(pairs) if pairs else "All Pairs",
        "window_hours": window_hours,
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M:%S"),
        "row_count": row_count,
        "failed_count": failed_count,
    }


def build_samples_email_context(
    connectors: list[str],
    pairs: list[str],
    row_count: int,
) -> dict:
    """Build the variable context used to render the spread samples email template."""
    now = datetime.datetime.now()
    return {
        "connectors": ", ".join(connectors) if connectors else "-",
        "pairs": ", ".join(pairs) if pairs else "-",
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M:%S"),
        "row_count": row_count,
    }


def parse_recipients(raw: str) -> tuple[list[str], list[str]]:
    """Split a comma/semicolon/whitespace separated string into valid and invalid emails."""
    if not raw:
        return [], []
    candidates = re.split(r"[,;\s]+", raw.strip())
    recipients = [c for c in candidates if c]
    valid = [c for c in recipients if EMAIL_RE.match(c)]
    invalid = [c for c in recipients if not EMAIL_RE.match(c)]
    return valid, invalid


def _xlsx_cell_value(value):
    """Coerce a pandas/numpy scalar into something openpyxl's write-only mode accepts."""
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def dataframes_to_xlsx_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Build an in-memory .xlsx file from a dict of {sheet_name: DataFrame}."""
    from openpyxl import Workbook

    export_sheets = sheets or {"Sheet1": pd.DataFrame()}
    workbook = Workbook(write_only=True)
    for sheet_name, df in export_sheets.items():
        safe_name = (sheet_name or "Sheet1")[:31]
        worksheet = workbook.create_sheet(title=safe_name)
        worksheet.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            worksheet.append([_xlsx_cell_value(v) for v in row])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.read()


def is_smtp_configured() -> bool:
    return bool(SMTP_HOST and SMTP_USERNAME and SMTP_PASSWORD)


MAX_EMAIL_ATTACHMENT_BYTES = 18 * 1024 * 1024


def send_email_with_xlsx(
    to_emails: list[str],
    subject: str,
    body: str,
    attachment_bytes: bytes,
    attachment_filename: str,
) -> None:
    """Send an email with an .xlsx attachment via SMTP. Raises on failure."""
    if not is_smtp_configured():
        raise RuntimeError(
            "SMTP is not configured. Set SMTP_HOST, SMTP_USERNAME and SMTP_PASSWORD in the .env file."
        )
    if not to_emails:
        raise ValueError("No valid recipient email addresses were provided.")
    if len(attachment_bytes) > MAX_EMAIL_ATTACHMENT_BYTES:
        size_mb = len(attachment_bytes) / (1024 * 1024)
        raise ValueError(
            f"The attachment is {size_mb:.1f} MB, which is too large to email reliably "
            "Use the Download button instead."
        )

    msg = MIMEMultipart()
    msg["From"] = SMTP_FROM_EMAIL or SMTP_USERNAME
    msg["To"] = ", ".join(to_emails)
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain"))

    part = MIMEApplication(
        attachment_bytes,
        _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    part.add_header("Content-Disposition", "attachment", filename=attachment_filename)
    msg.attach(part)

    smtp_class = smtplib.SMTP_SSL if SMTP_USE_TLS and SMTP_PORT == 465 else smtplib.SMTP
    with smtp_class(SMTP_HOST, SMTP_PORT, timeout=SMTP_TIMEOUT) as server:
        if SMTP_USE_TLS and smtp_class is smtplib.SMTP:
            server.starttls()
        server.login(SMTP_USERNAME, SMTP_PASSWORD)
        server.sendmail(SMTP_FROM_EMAIL or SMTP_USERNAME, to_emails, msg.as_string())
